#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
tasyrel_atb.py — converte o relatório "Utilização de Materiais/Medicamentos"
do TasyRel (.XLS BIFF2) na matriz da planilha de consumo de ATB da farmácia.

Uso:
    python3 tasyrel_atb.py atb_julho.XLS --planilha ATB2023.xlsx
    python3 tasyrel_atb.py atb_julho.XLS --planilha ATB2023.xlsx --escrever

Saídas:
    consumo_atb_<MES>_<ANO>.json   matriz + auditoria
    ATB2023_atualizada.xlsx        (apenas com --escrever) nova aba clonando o layout

Dependências: openpyxl (só para ler/escrever o .xlsx). O leitor do .XLS é stdlib.
"""

import json
import struct
import sys
import unicodedata
from argparse import ArgumentParser
from pathlib import Path

# ─────────────────────────────────────────────────────────────────────────────
# 1. Leitor BIFF2 (formato antigo que o TasyRel emite como ".XLS")
#    Registros usados: 0x0004 LABEL (texto), 0x0003 NUMBER (float),
#                      0x0002 INTEGER, 0x0001 BLANK
#    Header do registro: opcode (2 bytes LE) + length (2 bytes LE)
# ─────────────────────────────────────────────────────────────────────────────

def ler_biff2(caminho, encoding="cp1252"):
    """Devolve dict {(linha, coluna): valor} lendo o stream de registros BIFF2."""
    dados = Path(caminho).read_bytes()
    celulas = {}
    pos, fim = 0, len(dados)
    while pos + 4 <= fim:
        opcode, tam = struct.unpack_from("<HH", dados, pos)
        pos += 4
        if pos + tam > fim:
            break
        corpo = dados[pos:pos + tam]
        pos += tam

        if opcode == 0x0004 and tam >= 8:                 # LABEL
            lin, col = struct.unpack_from("<HH", corpo, 0)
            n = corpo[7]
            celulas[(lin, col)] = corpo[8:8 + n].decode(encoding, "replace").strip()
        elif opcode == 0x0003 and tam >= 15:              # NUMBER (IEEE 754)
            lin, col = struct.unpack_from("<HH", corpo, 0)
            celulas[(lin, col)] = struct.unpack_from("<d", corpo, 7)[0]
        elif opcode == 0x0002 and tam >= 9:               # INTEGER
            lin, col = struct.unpack_from("<HH", corpo, 0)
            celulas[(lin, col)] = float(struct.unpack_from("<H", corpo, 7)[0])
    return celulas


def _ler_xls_fallback(caminho):
    """Se um dia o TasyRel passar a exportar BIFF5/8, usa xlrd se disponível."""
    import xlrd  # noqa
    livro = xlrd.open_workbook(caminho, encoding_override="cp1252")
    aba = livro.sheet_by_index(0)
    return {(r, c): aba.cell_value(r, c)
            for r in range(aba.nrows) for c in range(aba.ncols)
            if aba.cell_value(r, c) not in ("", None)}


# ─────────────────────────────────────────────────────────────────────────────
# 2. Extração das linhas do relatório
# ─────────────────────────────────────────────────────────────────────────────

MESES = ["JANEIRO", "FEVEREIRO", "MARÇO", "ABRIL", "MAIO", "JUNHO",
         "JULHO", "AGOSTO", "SETEMBRO", "OUTUBRO", "NOVEMBRO", "DEZEMBRO"]

COL_SETOR, COL_MATERIAL, COL_QTDE, COL_VALOR = 1, 2, 9, 10


def norm(txt):
    """Normaliza para comparação: sem acento, maiúsculo, espaços colapsados."""
    if txt is None:
        return ""
    txt = unicodedata.normalize("NFKD", str(txt))
    txt = "".join(c for c in txt if not unicodedata.combining(c))
    return " ".join(txt.upper().split())


def ler_relatorio(caminho):
    """Devolve (lançamentos, cabeçalho). Cada lançamento: setor, cod, nome, qtde, valor."""
    try:
        cel = ler_biff2(caminho)
    except Exception:
        cel = _ler_xls_fallback(caminho)

    nlin = max(l for l, _ in cel) + 1
    periodo = next((v for (l, c), v in cel.items()
                    if isinstance(v, str) and v.startswith("Período")), "")
    mes = ano = None
    if "/" in periodo:
        partes = periodo.replace("Período de ", "").split()[0].split("/")
        if len(partes) == 3:
            mes, ano = MESES[int(partes[1]) - 1], partes[2]

    lancamentos, totais_setor = [], []
    for lin in range(nlin):
        setor = cel.get((lin, COL_SETOR))
        material = cel.get((lin, COL_MATERIAL))
        rotulo = cel.get((lin, 6))
        if isinstance(rotulo, str) and rotulo.startswith("Total"):
            totais_setor.append((rotulo, cel.get((lin, 9)), cel.get((lin, 10))))
            continue
        if not isinstance(setor, str) or not isinstance(material, str):
            continue
        if not setor or not material or norm(setor) == "SETOR":
            continue
        cod, _, nome = material.partition(" - ")
        lancamentos.append({
            "setor": setor.strip(),
            "cod": cod.strip(),
            "nome": nome.strip(),          # ATENÇÃO: truncado em ~25 caracteres
            "qtde": float(cel.get((lin, COL_QTDE)) or 0),
            "valor": float(cel.get((lin, COL_VALOR)) or 0),
        })

    total_geral = next((q for r, q, v in totais_setor if r.startswith("Total Geral")), None)
    return lancamentos, {"periodo": periodo, "mes": mes, "ano": ano,
                         "total_geral_qtde": total_geral}


# ─────────────────────────────────────────────────────────────────────────────
# 3. Layout da planilha da farmácia (lido da aba mais recente — não hardcoded)
# ─────────────────────────────────────────────────────────────────────────────

def ler_layout(ws):
    """Devolve (colunas, linhas, blocos) a partir de uma aba-modelo.

    colunas: {nome_normalizado_da_unidade: índice_da_coluna}
    linhas:  {(antimicrobiano_norm, apresentacao_norm): índice_da_linha}
    """
    colunas, linhas, blocos = {}, {}, []
    lin_cab = None
    for lin in range(1, ws.max_row + 1):
        a = ws.cell(lin, 1).value
        if isinstance(a, str) and norm(a) in ("ANTIMICROBIANO POR NOME GENERICO", "DEMAIS ATM"):
            lin_cab = lin
            for col in range(3, 20):
                v = ws.cell(lin_cab, col).value
                if isinstance(v, str) and v.strip():
                    colunas.setdefault(norm(v), col)
                    colunas.setdefault(v.strip(), col)
            blocos.append(lin_cab)
            continue
        if lin_cab and isinstance(a, str) and a.strip():
            b = ws.cell(lin, 2).value
            linhas[(norm(a), norm(b))] = lin
    return colunas, linhas, blocos


# ─────────────────────────────────────────────────────────────────────────────
# 4. Consolidação
# ─────────────────────────────────────────────────────────────────────────────

def consolidar(lancamentos, mapa, colunas, linhas):
    matriz, auditoria = {}, {
        "setores_nao_mapeados": {},
        "materiais_nao_mapeados": {},
        "destinos_inexistentes": [],
        "materiais_ignorados": {},
        "qtde_total_relatorio": 0.0,
        "qtde_total_alocada": 0.0,
        "qtde_ignorada": 0.0,
        "qtde_perdida": 0.0,
    }
    m_set, m_mat = mapa["setores"], mapa["materiais"]
    m_set_n = {norm(k): v for k, v in m_set.items()}

    for lc in lancamentos:
        q = lc["qtde"]
        auditoria["qtde_total_relatorio"] += q

        col_nome = m_set_n.get(norm(lc["setor"]), "__AUSENTE__")
        if col_nome == "__AUSENTE__":
            d = auditoria["setores_nao_mapeados"].setdefault(lc["setor"], 0)
            auditoria["setores_nao_mapeados"][lc["setor"]] = d + q
            auditoria["qtde_perdida"] += q
            continue

        if lc["cod"] not in m_mat:
            d = auditoria["materiais_nao_mapeados"].setdefault(
                f'{lc["cod"]} - {lc["nome"]}', 0)
            auditoria["materiais_nao_mapeados"][f'{lc["cod"]} - {lc["nome"]}'] = d + q
            auditoria["qtde_perdida"] += q
            continue

        destino = m_mat[lc["cod"]]
        if destino is None or col_nome is None:
            k = f'{lc["cod"]} - {lc["nome"]}'
            auditoria["materiais_ignorados"][k] = auditoria["materiais_ignorados"].get(k, 0) + q
            auditoria["qtde_ignorada"] += q
            continue

        chave_lin = (norm(destino[0]), norm(destino[1]))
        if chave_lin not in linhas or norm(col_nome) not in colunas:
            auditoria["destinos_inexistentes"].append(
                {"cod": lc["cod"], "destino": destino, "coluna": col_nome, "qtde": q})
            auditoria["qtde_perdida"] += q
            continue

        chave = (destino[0], destino[1], col_nome)
        matriz[chave] = matriz.get(chave, 0) + q
        auditoria["qtde_total_alocada"] += q

    return matriz, auditoria


# ─────────────────────────────────────────────────────────────────────────────
# 5. Escrita na planilha
# ─────────────────────────────────────────────────────────────────────────────

def escrever_aba(wb, aba_modelo, nome_nova_aba, matriz, colunas, linhas, mes, ano):
    nova = wb.copy_worksheet(aba_modelo)
    nova.title = nome_nova_aba

    for lin in linhas.values():                     # limpa a área de dados
        for col in range(3, 18):
            nova.cell(lin, col).value = None

    for lin in range(1, 4):                         # atualiza o cabeçalho
        for col in range(1, 6):
            v = nova.cell(lin, col).value
            if isinstance(v, str) and "REFERÊNCIA" in v.upper():
                if "ANO" in v.upper():
                    nova.cell(lin, col).value = f"ANO DE REFERÊNCIA: {ano}"
                else:
                    nova.cell(lin, col).value = f"MÊS DE REFERÊNCIA: {mes}"

    for (antimic, apres, unid), q in matriz.items():
        lin = linhas[(norm(antimic), norm(apres))]
        col = colunas[norm(unid)]
        nova.cell(lin, col).value = int(q) if float(q).is_integer() else q
    return nova


# ─────────────────────────────────────────────────────────────────────────────
# 6. CLI
# ─────────────────────────────────────────────────────────────────────────────

def main():
    ap = ArgumentParser(description="TasyRel .XLS → planilha de consumo de ATB")
    ap.add_argument("xls", help="relatório do TasyRel")
    ap.add_argument("--planilha", required=True, help="ATB2023.xlsx (modelo de layout)")
    ap.add_argument("--mapa", default=str(Path(__file__).with_name("mapa_atb.json")))
    ap.add_argument("--aba-modelo", default=None, help="aba usada como modelo (padrão: última)")
    ap.add_argument("--escrever", action="store_true", help="grava a nova aba no .xlsx")
    ap.add_argument("--saida", default=None)
    args = ap.parse_args()

    import openpyxl

    mapa = json.loads(Path(args.mapa).read_text(encoding="utf-8"))
    lancamentos, cab = ler_relatorio(args.xls)

    wb = openpyxl.load_workbook(args.planilha)
    modelo = wb[args.aba_modelo] if args.aba_modelo else wb[wb.sheetnames[-1]]
    colunas, linhas, _ = ler_layout(modelo)

    matriz, aud = consolidar(lancamentos, mapa, colunas, linhas)
    mes, ano = cab["mes"], cab["ano"]
    nome_aba = f"{mes} {ano}"

    # ── relatório de auditoria no terminal ────────────────────────────────────
    print(f"\n{cab['periodo']}   →   aba '{nome_aba}'")
    print(f"lançamentos lidos: {len(lancamentos)}   "
          f"modelo de layout: '{modelo.title}'")
    print(f"\nQuantidades  relatório={aud['qtde_total_relatorio']:.0f}  "
          f"alocada={aud['qtde_total_alocada']:.0f}  "
          f"ignorada={aud['qtde_ignorada']:.0f}  "
          f"PERDIDA={aud['qtde_perdida']:.0f}")
    if cab["total_geral_qtde"]:
        print(f"Total Geral do TasyRel: {cab['total_geral_qtde']:.0f} "
              f"(confere: {abs(cab['total_geral_qtde'] - aud['qtde_total_relatorio']) < 0.5})")

    for rotulo, chave in [("SETORES NÃO MAPEADOS", "setores_nao_mapeados"),
                          ("MATERIAIS NÃO MAPEADOS", "materiais_nao_mapeados"),
                          ("IGNORADOS POR REGRA (null)", "materiais_ignorados")]:
        if aud[chave]:
            print(f"\n{rotulo}:")
            for k, v in sorted(aud[chave].items(), key=lambda x: -x[1]):
                print(f"   {v:8.0f}  {k}")
    if aud["destinos_inexistentes"]:
        print("\nDESTINOS INEXISTENTES NA PLANILHA (corrigir mapa_atb.json):")
        for d in aud["destinos_inexistentes"]:
            print(f"   {d['qtde']:8.0f}  {d['cod']} → {d['destino']} / {d['coluna']}")

    # ── JSON ─────────────────────────────────────────────────────────────────
    saida_json = args.saida or f"consumo_atb_{(mes or 'MES').lower()}_{ano or ''}.json"
    aud["destinos_inexistentes"] = aud["destinos_inexistentes"]
    Path(saida_json).write_text(json.dumps({
        "periodo": cab["periodo"], "mes": mes, "ano": ano,
        "matriz": [{"antimicrobiano": a, "apresentacao": p, "unidade": u, "qtde": q}
                   for (a, p, u), q in sorted(matriz.items())],
        "auditoria": aud,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nJSON: {saida_json}")

    # ── XLSX ─────────────────────────────────────────────────────────────────
    if args.escrever:
        if nome_aba in wb.sheetnames:
            del wb[nome_aba]
        escrever_aba(wb, modelo, nome_aba, matriz, colunas, linhas, mes, ano)
        destino = Path(args.planilha).with_name(
            Path(args.planilha).stem + "_atualizada.xlsx")
        wb.save(destino)
        print(f"XLSX: {destino}")


if __name__ == "__main__":
    main()
