"""
cve_writer.py — orquestrador: lê os parsers, escreve SÓ inputs no CVE (nunca fórmula)
e devolve o diff de revisão. Não recalcula (o Excel recalcula ao abrir).

build_cve(controles_path, cve_master_path, out_path, meses, numeradores=None) -> list[diff]
build_cve_dialise(setor_path, dial_master_path, out_path, meses, hd_num=None) -> list[diff]
  diff row = [mes, aba, celula, valor_antigo, valor_novo, fonte, status]
  status ∈ {OK, DIVERGE, NOVO, PENDENTE}
"""
import shutil
import openpyxl
from husf_pipeline.parsers import cve_controles as CC
from husf_pipeline.parsers import cve_dialise as CD

def _anchor(ws, mes):
    for r in range(1, ws.max_row+1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v.strip() == mes:
            return r
    return None

def _write_numeradores(CVE, mes, num, put):
    a = _anchor(CVE['Plan2 - UTI'], mes)
    if a:
        for u, off in [('uti1',3),('uti2',4),('utipe',8)]:
            d = num['plan2'][u]
            put('Plan2 - UTI', a+off, 2, d['pn'],   f'SURA/{u}')
            put('Plan2 - UTI', a+off, 3, d['ipcs'], f'SURA/{u}')
            put('Plan2 - UTI', a+off, 4, d['iu'],   f'SURA/{u}')
    a3 = _anchor(CVE['Plan3 - UTINeo'], mes)
    if a3:
        for i, b in enumerate('ABCDE'):
            d = num['plan3'][b]
            put('Plan3 - UTINeo', a3+3+i, 2, d['pn'],        'SURA/Neo')
            put('Plan3 - UTINeo', a3+3+i, 3, d['ipcs_lab'],  'SURA/Neo')
            put('Plan3 - UTINeo', a3+3+i, 4, d['ipcs_clin'], 'SURA/Neo')
    a1 = _anchor(CVE['Plan1 - ISC Especialidade'], mes)
    if a1:
        ws = CVE['Plan1 - ISC Especialidade']
        for r in range(a1+2, a1+13):
            code = ws.cell(row=r, column=1).value
            if isinstance(code, str) and code.strip() in num['plan1']:
                put('Plan1 - ISC Especialidade', r, 2, num['plan1'][code.strip()], 'SURA/ISC')
    a1b = _anchor(CVE['Plan1B - ISC Cirurgia'], mes)
    if a1b:
        ws = CVE['Plan1B - ISC Cirurgia']
        for r in range(a1b+2, a1b+15):
            typ = ws.cell(row=r, column=1).value
            if isinstance(typ, str) and typ.strip() in num['plan1b']:
                d = num['plan1b'][typ.strip()]
                put('Plan1B - ISC Cirurgia', r, 2, d['incisional'], 'SURA/ISC')
                if d.get('cavidade'):
                    put('Plan1B - ISC Cirurgia', r, 3, d['cavidade'], 'SURA/ISC')

def _write_month(CVE, OLD, S, mes, diff, numeradores=None):
    mi = CC.MESES.index(mes)
    dU  = CC.dados_uti(S, mi)
    left, right = CC.cirurgias(S, mes)
    atb = CC.consumo_atb(S, mes)
    chk = CC.checklist_cvc(S, mes)
    alc = CC.alcool_sabonete(S, mi)

    def put(sheet, r, c, val, fonte):
        if val is None: return
        cell = CVE[sheet].cell(row=r, column=c)
        old = OLD[sheet].cell(row=r, column=c).value
        oldn = CC.num(old)
        if oldn is not None and abs(oldn - val) < 0.5: status = 'OK'
        elif oldn is not None:                          status = 'DIVERGE'
        else:                                           status = 'NOVO'
        diff.append([mes, sheet, cell.coordinate, old, val, fonte, status])
        cell.value = val

    # Plan2 device-days + processo J/K/L
    a = _anchor(CVE['Plan2 - UTI'], mes)
    if a:
        for key, off in [('uti1',3),('uti2',4),('utipe',8)]:
            d = dU[key]
            put('Plan2 - UTI', a+off, 5, d['vm'], f'Dados UTI/{key}')
            put('Plan2 - UTI', a+off, 6, d['ct'], f'Dados UTI/{key}')
            put('Plan2 - UTI', a+off, 7, d['sv'], f'Dados UTI/{key}')
            put('Plan2 - UTI', a+off, 8, d['pd'], f'Dados UTI/{key}')
        for u, off in [('uti1',3),('uti2',4)]:
            vals = chk.get(u)
            if vals:
                for j, col in enumerate((10,11,12)):
                    put('Plan2 - UTI', a+off, col, vals[j], f'Checklist CVC/{u}')
    # Plan3 Neo device-days
    a3 = _anchor(CVE['Plan3 - UTINeo'], mes)
    if a3:
        for i, b in enumerate('ABCDE'):
            d = dU['neo'][b]
            put('Plan3 - UTINeo', a3+3+i, 5, d['vm'], 'Dados UTI/Neo')
            put('Plan3 - UTINeo', a3+3+i, 6, d['ct'], 'Dados UTI/Neo')
            put('Plan3 - UTINeo', a3+3+i, 7, d['pd'], 'Dados UTI/Neo')
    # Plan1 CL
    a1 = _anchor(CVE['Plan1 - ISC Especialidade'], mes)
    if a1:
        ws = CVE['Plan1 - ISC Especialidade']
        for r in range(a1+2, a1+13):
            code = ws.cell(row=r, column=1).value
            if isinstance(code, str) and code.strip() and code.strip() != 'Total':
                k = CC.norm(code); k = CC.CODE_ALIAS.get(k, k)
                put('Plan1 - ISC Especialidade', r, 3, left.get(k), 'Cirurgias limpas')
    # Plan1B total
    a1b = _anchor(CVE['Plan1B - ISC Cirurgia'], mes)
    if a1b:
        ws = CVE['Plan1B - ISC Cirurgia']
        for r in range(a1b+2, a1b+15):
            typ = ws.cell(row=r, column=1).value
            if isinstance(typ, str) and typ.strip() and typ.strip() != 'Total':
                put('Plan1B - ISC Cirurgia', r, 4, right.get(CC.norm(typ)), 'Cirurgias limpas')
    # Plan6 nº unidades UTIA
    a6 = _anchor(CVE['Plan6 - DDD'], mes)
    if a6:
        ws = CVE['Plan6 - DDD']
        for r in range(a6+2, a6+180):
            g = ws.cell(row=r, column=1).value
            if isinstance(g, str) and g.strip() in CC.MESES: break
            ap = ws.cell(row=r, column=2).value
            if isinstance(g,str) and g.strip() and isinstance(ap,str) and ap.strip():
                val = atb.get((CC.gkey(g), CC.norm(ap)))
                if val is not None:
                    put('Plan6 - DDD', r, 3, val, 'Consumo mensal ATB (N+O)')
    # Plan4 álcool/sabonete (D=pac-dia é fórmula)
    r1, r2 = 16+mi, 32+mi
    put('Plan4 - P.alcoolico e sabonete', r1, 2, alc['ad_alc'],  'Dados UTI/frascos×500')
    put('Plan4 - P.alcoolico e sabonete', r1, 3, alc['ad_sab'],  'Dados UTI/frascos×500')
    put('Plan4 - P.alcoolico e sabonete', r2, 2, alc['ped_alc'], 'Dados UTI/frascos×500')
    put('Plan4 - P.alcoolico e sabonete', r2, 3, alc['ped_sab'], 'Dados UTI/frascos×500')
    put('Plan4 - P.alcoolico e sabonete', r2, 7, alc['neo_alc'], 'Dados UTI/frascos×500')
    put('Plan4 - P.alcoolico e sabonete', r2, 8, alc['neo_sab'], 'Dados UTI/frascos×500')
    # Numeradores SURA (opcional) + pendências
    if numeradores and mi+1 in numeradores:
        nm = numeradores[mi+1]
        _write_numeradores(CVE, mes, nm, put)
        for p in nm.get('plan1_pendente', []):
            diff.append([mes, 'Plan1 - ISC Especialidade', '(revisar)', '', p['n'],
                         f"SURA/ISC etiol={p['etiol']} equipe={p['equipe']} cir={p['cirurgia']}", 'PENDENTE'])

def build_cve(controles_path, cve_master_path, out_path, meses, numeradores=None):
    shutil.copy(cve_master_path, out_path)
    S   = openpyxl.load_workbook(controles_path, data_only=True)
    OLD = openpyxl.load_workbook(cve_master_path, data_only=True)
    CVE = openpyxl.load_workbook(out_path)   # preserva fórmulas
    diff = []
    for mes in meses:
        _write_month(CVE, OLD, S, mes, diff, numeradores)
    CVE.save(out_path)
    return diff

# ── CVE de diálise (destino separado) ──────────────────────────────────────────
# aba Hemodiálise, linha do mês = 32 + mi. Denom/extra ← setor; numeradores ← SURA.
_DIAL_SETOR = {'temp':2,'perm':3,'fist':4,'hosp':6,'cvc3m':8,'soro':10,'hcvneg':11,'obito':13,'vanco':31}
_DIAL_SURA  = {('ilav','temp'):15, ('ilav','perm'):18, ('ilav','fist'):21,   # O,R,U
               ('bact','temp'):23, ('bact','perm'):26, ('bact','fist'):29}    # W,Z,AC

def build_cve_dialise(setor_path, dial_master_path, out_path, meses, hd_num=None):
    """Preenche o CVE de diálise: denom/contagens do setor + numeradores IRAS do SURA."""
    shutil.copy(dial_master_path, out_path)
    OLD = openpyxl.load_workbook(dial_master_path, data_only=True)
    CVE = openpyxl.load_workbook(out_path)
    ws, wso = CVE['Hemodiálise'], OLD['Hemodiálise']
    diff = []
    for mes in meses:
        mi = CC.MESES.index(mes)
        row = 32 + mi
        def put(col, val, fonte):
            if val is None: return
            old = wso.cell(row=row, column=col).value
            oldn = CC.num(old)
            if oldn is not None and abs(oldn - val) < 0.5: st = 'OK'
            elif oldn is not None:                          st = 'DIVERGE'
            else:                                           st = 'NOVO'
            diff.append([mes, 'Hemodiálise', ws.cell(row=row, column=col).coordinate, old, val, fonte, st])
            ws.cell(row=row, column=col).value = val
        setor = CD.parse_hemodialise(setor_path, mi)
        for k, col in _DIAL_SETOR.items():
            put(col, CC.num(setor.get(k)), 'setor/hemodiálise')
        if hd_num and (mi+1) in hd_num:
            h = hd_num[mi+1]
            for (grupo, acesso), col in _DIAL_SURA.items():
                put(col, h[grupo][acesso], f'SURA/{grupo}_{acesso}')
            for p in h.get('pendente', []):
                diff.append([mes, 'Hemodiálise', '(revisar)', '', p['n'],
                             f"SURA HD sem acesso: {p['iras']}", 'PENDENTE'])
    CVE.save(out_path)
    return diff
