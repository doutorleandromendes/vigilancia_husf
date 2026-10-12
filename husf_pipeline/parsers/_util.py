"""
Utilitários compartilhados dos parsers + abstração de "provedor de valores".

A ideia central: um parser nunca abre um arquivo nem chama a API do Google.
Ele recebe um `provider` — um objeto que, dado o nome de uma aba, devolve a
matriz de valores daquela aba. Assim o MESMO parser roda contra um .xlsx
(para teste/desenvolvimento) ou contra a Google Sheets API (em produção),
sem nenhuma alteração na lógica de parsing.

    provider.values("ABR")  ->  [[c0, c1, ...], [c0, c1, ...], ...]
    provider.tabs()         ->  ["JAN", "FEV", ...]
"""

import unicodedata

# ── Constantes de mês ────────────────────────────────────────────────────────
MESES_EN = ['January', 'February', 'March', 'April', 'May', 'June',
            'July', 'August', 'September', 'October', 'November', 'December']
MESES_PT_ABBR = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun',
                 'jul', 'ago', 'set', 'out', 'nov', 'dez']
MESES_PT_FULL = ['Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho',
                 'Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro']


def mes_label(mes_idx, ano):
    """(3, 2026) -> 'abr/26'"""
    return f"{MESES_PT_ABBR[mes_idx]}/{str(ano)[-2:]}"


def parse_mes_arg(s):
    """'abr/26' -> (3, 2026).  Aceita também 'abr/2026'."""
    abbr, ano = s.lower().split('/')
    idx = MESES_PT_ABBR.index(abbr.strip())
    ano = int(ano)
    if ano < 100:
        ano += 2000
    return idx, ano


# ── Normalização / conversão numérica ────────────────────────────────────────
def norm(s):
    """Remove acentos, faz upper e strip — para comparar rótulos com robustez."""
    if s is None:
        return ''
    s = str(s).strip()
    s = ''.join(c for c in unicodedata.normalize('NFD', s)
                if unicodedata.category(c) != 'Mn')
    return s.upper()


_ERR_MARKERS = ('#DIV/0!', '#N/A', '#VALUE!', '#REF!', '#NAME?', '#NULL!', '#NUM!')


def num(v, on_error=0.0):
    """
    Converte célula em float. Diferente do safeN do GAS: células de ERRO do
    Excel (#DIV/0! etc.) retornam `on_error` — por padrão 0.0, mas os parsers
    passam None onde um zero falso seria perigoso (ex.: mês sem dado nos DOTs).
    """
    if v is None or v == '':
        return on_error
    if isinstance(v, str):
        sv = v.strip()
        if sv.upper() in _ERR_MARKERS or (sv and sv[0] == '#'):
            return on_error
        sv = sv.replace(',', '.')
        try:
            return float(sv)
        except ValueError:
            return on_error
    try:
        return float(v)
    except (TypeError, ValueError):
        return on_error


def r2(x):
    """Arredonda em 2 casas, preservando None."""
    return None if x is None else round(x, 2)


def r1(x):
    return None if x is None else round(x, 1)


def calc_di(n, den):
    """Densidade de incidência: n / dispositivo-dia * 1000 (2 casas)."""
    n = num(n)
    den = num(den)
    if not den:
        return 0.0
    return round(n / den * 1000, 2)


def calc_tx(n, den):
    """Taxa de utilização: dispositivo-dia / pac-dia * 100 (1 casa)."""
    n = num(n)
    den = num(den)
    if not den:
        return 0.0
    return round(n / den * 100, 1)


# ── Provedores de valores ────────────────────────────────────────────────────
class XlsxProvider:
    """Lê abas de um arquivo .xlsx local (desenvolvimento/teste)."""

    def __init__(self, path):
        from openpyxl import load_workbook
        self._wb = load_workbook(path, read_only=True, data_only=True)
        # cache por aba (matrizes podem ser grandes)
        self._cache = {}

    def tabs(self):
        return list(self._wb.sheetnames)

    def values(self, tab):
        if tab not in self._cache:
            ws = self._wb[tab]
            self._cache[tab] = [list(r) for r in ws.iter_rows(values_only=True)]
        return self._cache[tab]


class GSheetProvider:
    """
    Lê abas de uma planilha do Google via API (produção).
    Mesma interface do XlsxProvider — os parsers não percebem a diferença.

    Requer `gspread` e uma service account. Importado preguiçosamente para que
    o caminho de teste (xlsx) funcione sem essas dependências instaladas.
    """

    def __init__(self, spreadsheet_id, creds_path):
        import gspread
        gc = gspread.service_account(filename=creds_path)
        self._ss = gc.open_by_key(spreadsheet_id)
        self._cache = {}

    def tabs(self):
        return [w.title for w in self._ss.worksheets()]

    def values(self, tab):
        if tab not in self._cache:
            self._cache[tab] = self._ss.worksheet(tab).get_all_values()
        return self._cache[tab]


# ── Seleção de aba do mês (tolerante a EN / PT-abbr / PT-full / número) ───────
def find_month_tab(provider, mes_idx):
    """
    Devolve o nome real da aba do mês, tentando, em ordem e case-insensitive:
    inglês, PT-abreviado, PT-completo, número (1..12).
    Levanta erro se nenhuma casar.
    """
    candidates = [
        MESES_EN[mes_idx],
        MESES_PT_ABBR[mes_idx],
        MESES_PT_FULL[mes_idx],
        str(mes_idx + 1),
    ]
    tabs = provider.tabs()
    tab_norm = {norm(t): t for t in tabs}
    for c in candidates:
        if norm(c) in tab_norm:
            return tab_norm[norm(c)]
    raise ValueError(
        f"Aba do mês não encontrada (tentou {candidates}). Abas: {tabs}"
    )
