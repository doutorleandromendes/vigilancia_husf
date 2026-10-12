"""
parse_dialise_hist.py — extrai série mensal de ILAV e Bacteremia
das planilhas CVE Hemodiálise anuais (arquivos de dezembro).

Uso:
  python3 parse_dialise_hist.py \
    --2023 PLANILHA_CVE_dialise_HUSF_122023.xlsx \
    --2024 PLANILHA_CVE_dialise_HUSF_122024.xlsx \
    --2025 PLANILHA_CVE_dialise_HUSF_122025.xlsx \
    --json  data_iras.json

Colunas extraídas (consistentes em 2023/2024/2025):
  col 16 = ILAV cateter temporário  /1000 cat-dia  → hd.ilavTemp
  col 19 = ILAV cateter permanente  /1000 cat-dia  → hd.ilavPerm
  col 24 = Bacteremia cateter temp  /1000 cat-dia  → hd.bactTemp
  col 27 = Bacteremia cateter perm  /1000 cat-dia  → hd.bactPerm
"""

import json, argparse
from pathlib import Path
from openpyxl import load_workbook

ABBR = ['jan','fev','mar','abr','mai','jun','jul','ago','set','out','nov','dez']
MESES_PT = ['Janeiro','Fevereiro','Março','Abril','Maio','Junho',
            'Julho','Agosto','Setembro','Outubro','Novembro','Dezembro']

# Offset da linha de Janeiro para cada ano
DATA_START = {2023: 37, 2024: 31, 2025: 31, 2026: 31}

COLS = {
    'ilavTemp': 16,
    'ilavPerm': 19,
    'bactTemp': 24,
    'bactPerm': 27,
}

def num(v):
    if v is None: return None
    try: f = float(str(v)); return None if f == 0 else round(f, 4)
    except: return None

def parse_file(path, ano):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb['Hemodiálise']
    rows = list(ws.iter_rows(values_only=True))
    start = DATA_START.get(ano, 31)
    out = {k: [] for k in COLS}
    for mi in range(12):
        ri = start + mi
        if ri >= len(rows): break
        row = rows[ri]
        if not row or not row[0]: continue
        label = f"{ABBR[mi]}/{str(ano)[-2:]}"
        for key, ci in COLS.items():
            v = num(row[ci]) if ci < len(row) else None
            out[key].append({'p': label, 'v': v if v is not None else 0})
    return out

def merge_series(existing, new_pts):
    """Insere pontos históricos sem sobrescrever os existentes."""
    existing_labels = {p['p'] for p in existing}
    added = [p for p in new_pts if p['p'] not in existing_labels]
    MO = {a: i+1 for i, a in enumerate(ABBR)}
    def lk(p):
        m, y = p['p'].split('/')
        return (2000+int(y), MO[m])
    combined = sorted(added + existing, key=lk)
    return combined, len(added)

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--2023', dest='f2023')
    ap.add_argument('--2024', dest='f2024')
    ap.add_argument('--2025', dest='f2025')
    ap.add_argument('--json', dest='json_path', default='data_iras.json')
    args = ap.parse_args()

    d = json.load(open(args.json_path, encoding='utf-8'))
    d.setdefault('hd', {})

    total_added = 0
    for ano, fpath in [(2023, args.f2023), (2024, args.f2024), (2025, args.f2025)]:
        if not fpath: continue
        print(f"\n→ {ano}: {fpath}")
        parsed = parse_file(fpath, ano)
        for key, new_pts in parsed.items():
            d['hd'].setdefault(key, [])
            d['hd'][key], n = merge_series(d['hd'][key], new_pts)
            total_added += n
            print(f"  {key}: +{n} pontos (total {len(d['hd'][key])})")

    json.dump(d, open(args.json_path, 'w', encoding='utf-8'),
              ensure_ascii=False, indent=2)
    print(f"\n✓ data_iras.json atualizado ({total_added} pontos adicionados)")

if __name__ == '__main__':
    main()
